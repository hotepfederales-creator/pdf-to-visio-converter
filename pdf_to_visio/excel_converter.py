"""Excel wire-list to Visio-compatible harness drawing conversion.

The module reads a spreadsheet of source/destination pins and renders a
Visio-importable SVG harness drawing. When Microsoft Visio and pywin32 are
available, the same SVG can be placed into a native .vsdx file.
"""

from __future__ import annotations

from dataclasses import dataclass
import html
import os
from pathlib import Path
import tempfile
from typing import Any, Iterable, Mapping, Sequence

_HEADER_ALIASES = {
    "source_connector": (
        "sourceconnector",
        "srcconnector",
        "fromconnector",
        "fromconn",
        "sourceconn",
    ),
    "source_pin": ("sourcepin", "srcpin", "frompin"),
    "destination_connector": (
        "destinationconnector",
        "destconnector",
        "dstconnector",
        "toconnector",
        "toconnect",
        "toconn",
    ),
    "destination_pin": ("destinationpin", "destpin", "dstpin", "topin"),
    "signal": ("signal", "net", "circuit", "function"),
    "wire_color": ("wirecolor", "color", "colour", "wirecolour"),
    "gauge": ("gauge", "awg", "wiregauge", "size"),
    "breakout": (
        "breakout",
        "breakoutjack",
        "breakoutpin",
        "testpoint",
        "testpointid",
    ),
}

_REQUIRED_COLUMNS = (
    "source_connector",
    "source_pin",
    "destination_connector",
    "destination_pin",
)

_COLOR_MAP = {
    "black": "#111827",
    "blk": "#111827",
    "white": "#f9fafb",
    "wht": "#f9fafb",
    "red": "#dc2626",
    "blue": "#2563eb",
    "green": "#16a34a",
    "yellow": "#facc15",
    "orange": "#f97316",
    "brown": "#92400e",
    "violet": "#7c3aed",
    "purple": "#7c3aed",
    "gray": "#6b7280",
    "grey": "#6b7280",
}


@dataclass(frozen=True)
class Endpoint:
    """A connector pin endpoint from a wire-list row."""

    connector: str
    pin: str


@dataclass(frozen=True)
class WireConnection:
    """One source-to-destination harness connection."""

    source: Endpoint
    destination: Endpoint
    signal: str = ""
    wire_color: str = ""
    gauge: str = ""
    breakout: str = ""


def load_connections_from_excel(
    excel_path: str,
    sheet: str | int | None = None,
) -> list[WireConnection]:
    """Read an Excel wire list into normalized harness connections.

    Required columns are source connector, source pin, destination connector,
    and destination pin. Header names are matched case-insensitively and may use
    common aliases like ``From Connector``, ``From Pin``, ``To Connector``, and
    ``To Pin``.
    """
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependent
        raise ImportError(
            "openpyxl is required for Excel wire-list conversion: pip install openpyxl"
        ) from exc

    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    worksheet = _select_sheet(workbook, sheet)

    rows = worksheet.iter_rows(values_only=True)
    try:
        header_row = next(rows)
    except StopIteration as exc:
        workbook.close()
        raise ValueError("Excel sheet is empty") from exc

    column_map = _map_headers(header_row)
    missing = [column for column in _REQUIRED_COLUMNS if column not in column_map]
    if missing:
        workbook.close()
        raise ValueError(
            "Missing required Excel columns: "
            f"{', '.join(missing)}. Found: {_format_found_headers(header_row)}"
        )

    connections: list[WireConnection] = []
    for row in rows:
        connection = _connection_from_row(row, column_map)
        if connection is not None:
            connections.append(connection)

    workbook.close()
    if not connections:
        raise ValueError("Excel sheet has headers but no wire-list rows")
    return connections


def convert_excel_to_visio(
    excel_path: str,
    output_path: str,
    sheet: str | int | None = None,
    title: str | None = None,
) -> str:
    """Convert an Excel wire list into a Visio-ready drawing.

    ``.svg`` outputs are importable into Visio. ``.vsdx`` outputs require
    Microsoft Visio plus ``pywin32`` on Windows; the converter creates a new
    Visio document and imports the generated SVG onto the first page.
    """
    connections = load_connections_from_excel(excel_path, sheet=sheet)
    output_suffix = Path(output_path).suffix.lower()
    drawing_title = title or Path(excel_path).stem

    if output_suffix in {"", ".svg"}:
        return write_harness_svg(connections, output_path, title=drawing_title)

    if output_suffix == ".vsdx":
        return _write_vsdx_with_visio(connections, output_path, title=drawing_title)

    raise ValueError("Excel output path must end in .svg or .vsdx")


def write_harness_svg(
    connections: Sequence[WireConnection],
    output_path: str,
    title: str = "Harness Drawing",
) -> str:
    """Write a Visio-importable SVG harness drawing."""
    if not connections:
        raise ValueError("At least one wire connection is required")

    width = 980
    top = 92
    row_gap = 46
    height = max(260, top + len(connections) * row_gap + 80)
    has_breakout = any(connection.breakout for connection in connections)

    parts = [_svg_header(width, height, title)]
    parts.append(_title_block(title, width))
    parts.append(_legend(width))
    if has_breakout:
        parts.append(_breakout_box(width / 2 - 95, 50, 190, height - 92))

    for index, connection in enumerate(connections):
        y = top + index * row_gap
        parts.append(_render_connection_row(connection, y, has_breakout))

    parts.append("</svg>\n")

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)
    with open(output_path, "w", encoding="utf-8") as file:
        file.write("".join(parts))
    return os.path.abspath(output_path)


def _select_sheet(workbook: Any, sheet: str | int | None) -> Any:
    if sheet is None:
        return workbook.active
    if isinstance(sheet, int):
        try:
            return workbook.worksheets[sheet]
        except IndexError as exc:
            raise ValueError(f"Excel sheet index {sheet} is out of range") from exc
    try:
        return workbook[sheet]
    except KeyError as exc:
        raise ValueError(f"Excel sheet not found: {sheet}") from exc


def _normalize_header(value: object) -> str:
    return "".join(char for char in str(value or "").lower() if char.isalnum())


def _map_headers(header_row: Iterable[object]) -> dict[str, int]:
    normalized = [_normalize_header(header) for header in header_row]
    column_map: dict[str, int] = {}
    for canonical, aliases in _HEADER_ALIASES.items():
        for index, header in enumerate(normalized):
            if header in aliases:
                column_map[canonical] = index
                break
    return column_map


def _format_found_headers(header_row: Iterable[object]) -> str:
    headers = [str(header) for header in header_row if header not in (None, "")]
    return ", ".join(headers) if headers else "none"


def _connection_from_row(
    row: Sequence[object], column_map: Mapping[str, int]
) -> WireConnection | None:
    source_connector = _cell(row, column_map["source_connector"])
    source_pin = _cell(row, column_map["source_pin"])
    destination_connector = _cell(row, column_map["destination_connector"])
    destination_pin = _cell(row, column_map["destination_pin"])

    if not any((source_connector, source_pin, destination_connector, destination_pin)):
        return None

    missing = [
        name
        for name, value in (
            ("source_connector", source_connector),
            ("source_pin", source_pin),
            ("destination_connector", destination_connector),
            ("destination_pin", destination_pin),
        )
        if not value
    ]
    if missing:
        raise ValueError(f"Wire-list row is missing values for: {', '.join(missing)}")

    return WireConnection(
        source=Endpoint(source_connector, source_pin),
        destination=Endpoint(destination_connector, destination_pin),
        signal=_optional_cell(row, column_map, "signal"),
        wire_color=_optional_cell(row, column_map, "wire_color"),
        gauge=_optional_cell(row, column_map, "gauge"),
        breakout=_optional_cell(row, column_map, "breakout"),
    )


def _cell(row: Sequence[object], index: int) -> str:
    if index >= len(row) or row[index] is None:
        return ""
    return str(row[index]).strip()


def _optional_cell(
    row: Sequence[object], column_map: Mapping[str, int], column: str
) -> str:
    index = column_map.get(column)
    return "" if index is None else _cell(row, index)


def _wire_stroke(color_name: str) -> str:
    normalized = _normalize_header(color_name.split("/")[0].split("-")[0])
    return _COLOR_MAP.get(normalized, "#111827")


def _svg_header(width: int, height: int, title: str) -> str:
    escaped_title = html.escape(title)
    return (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" '
        f'viewBox="0 0 {width} {height}">\n'
        f"<title>{escaped_title}</title>\n"
        "<style>"
        ".box{fill:#f8fafc;stroke:#334155;stroke-width:2}"
        ".pin{fill:#fff;stroke:#0f172a;stroke-width:1.5}"
        ".wire{fill:none;stroke-width:3;stroke-linecap:round}"
        ".label{font-family:Segoe UI,Arial,sans-serif;font-size:13px;fill:#0f172a}"
        ".small{font-size:11px;fill:#475569}"
        ".title{font-size:24px;font-weight:700}"
        "</style>\n"
    )


def _title_block(title: str, width: int) -> str:
    return (
        f'<text x="24" y="34" class="label title">{html.escape(title)}</text>\n'
        f'<line x1="24" y1="48" x2="{width - 24}" y2="48" '
        'stroke="#cbd5e1" stroke-width="1"/>\n'
    )


def _legend(width: int) -> str:
    x = width - 310
    return (
        f'<rect x="{x}" y="14" width="286" height="26" rx="6" '
        'fill="#f1f5f9" stroke="#cbd5e1"/>\n'
        f'<text x="{x + 12}" y="32" class="label small">'
        "Excel wire list → harness / breakout diagram</text>\n"
    )


def _breakout_box(x: float, y: float, width: float, height: float) -> str:
    return (
        f'<rect x="{x:.1f}" y="{y:.1f}" width="{width:.1f}" height="{height:.1f}" '
        'rx="14" fill="#ecfeff" stroke="#0891b2" stroke-width="2"/>\n'
        f'<text x="{x + width / 2:.1f}" y="{y + 26:.1f}" '
        'text-anchor="middle" class="label">Breakout Box</text>\n'
    )


def _render_connection_row(
    connection: WireConnection, y: int, has_breakout: bool
) -> str:
    left_x = 58
    right_x = 758
    pin_radius = 8
    source = connection.source
    destination = connection.destination
    stroke = _wire_stroke(connection.wire_color)
    wire_label = _wire_label(connection)

    source_text = f"{source.connector} pin {source.pin}"
    destination_text = f"{destination.connector} pin {destination.pin}"
    parts = [
        f'<circle class="pin" cx="{left_x}" cy="{y}" r="{pin_radius}"/>\n',
        f'<text x="{left_x + 18}" y="{y - 7}" class="label">'
        f"{html.escape(source_text)}</text>\n",
        f'<text x="{left_x + 18}" y="{y + 10}" class="label small">'
        f"{html.escape(wire_label)}</text>\n",
        f'<circle class="pin" cx="{right_x}" cy="{y}" r="{pin_radius}"/>\n',
        f'<text x="{right_x + 18}" y="{y - 7}" class="label">'
        f"{html.escape(destination_text)}</text>\n",
    ]

    if has_breakout and connection.breakout:
        breakout_x = 490
        parts.extend(
            [
                f'<circle class="pin" cx="{breakout_x}" cy="{y}" r="{pin_radius}"/>\n',
                f'<text x="{breakout_x}" y="{y - 14}" text-anchor="middle" '
                f'class="label small">{html.escape(connection.breakout)}</text>\n',
                f'<path class="wire" d="M {left_x + pin_radius} {y} '
                f'L {breakout_x - pin_radius} {y}" stroke="{stroke}"/>\n',
                f'<path class="wire" d="M {breakout_x + pin_radius} {y} '
                f'L {right_x - pin_radius} {y}" stroke="{stroke}"/>\n',
            ]
        )
    else:
        parts.append(
            f'<path class="wire" d="M {left_x + pin_radius} {y} '
            f'L {right_x - pin_radius} {y}" stroke="{stroke}"/>\n'
        )

    return "".join(parts)


def _wire_label(connection: WireConnection) -> str:
    fields = [connection.signal, connection.wire_color, connection.gauge]
    label = " | ".join(field for field in fields if field)
    return label or "wire"


def _write_vsdx_with_visio(
    connections: Sequence[WireConnection], output_path: str, title: str
) -> str:
    try:
        import win32com.client  # type: ignore[import-not-found]
    except ImportError as exc:  # pragma: no cover - requires Windows/pywin32
        raise RuntimeError(
            "Native .vsdx export requires Microsoft Visio automation via pywin32. "
            "Install pywin32 and run on Windows with Visio installed, or export .svg."
        ) from exc

    out_dir = os.path.dirname(output_path)
    if out_dir:
        os.makedirs(out_dir, exist_ok=True)

    with tempfile.NamedTemporaryFile(suffix=".svg", delete=False) as tmp:
        tmp_svg = tmp.name
    try:
        write_harness_svg(connections, tmp_svg, title=title)
        visio = win32com.client.Dispatch("Visio.Application")
        visio.Visible = False
        document = visio.Documents.Add("")
        page = visio.ActivePage
        page.Import(tmp_svg)
        document.SaveAs(os.path.abspath(output_path))
        document.Close()
        visio.Quit()
    finally:
        if os.path.exists(tmp_svg):
            os.unlink(tmp_svg)

    return os.path.abspath(output_path)
